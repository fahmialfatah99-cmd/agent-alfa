"""Unit tests verifying bugfixes for:
1. Swarm auto-trigger prevention in single agent mode
2. Swarm cancel flag cleanup and session isolation
3. WhatsApp digest formatting preservation
4. wa_get_incoming noise filtering (@broadcast, @newsletter)
"""

import asyncio
import json
import os
import time
import pytest
from unittest.mock import patch, MagicMock, AsyncMock
from alfa.core.runtime_ctx import current_user_prompt_var, get_current_user_prompt
from alfa.tools.system.management import conduct_ai_meeting
from alfa.tools.system.integrations import wa_get_incoming
from alfa.swarm import workspace_hygiene
from alfa.bot.wa_digest import build_digest_text, summarize_records
from alfa.bot.pipeline.tool_config import prepare_tools_and_configs


@pytest.fixture
def isolated_db(tmp_path, monkeypatch):
    db_file = str(tmp_path / "test_bugfix.db")
    monkeypatch.setenv("ALFA_DB_PATH", db_file)
    monkeypatch.delenv("DATABASE_URL", raising=False)
    from alfa.core.db import connection as _conn

    _conn.init_db_sync()
    return db_file


def test_conduct_ai_meeting_single_mode_guard():
    # User prompt does not contain meeting keywords
    current_user_prompt_var.set("secanggih apa kamu?")
    with patch("alfa.bot.mode_commands.get_telegram_mode_sync", return_value="single"):
        res = conduct_ai_meeting(topic="Rapat kecanggihan sistem")
        assert res["status"] == "skipped"
        assert "DILARANG meluncurkan swarm" in res["message"]


def test_conduct_ai_meeting_single_mode_allowed_when_requested():
    current_user_prompt_var.set("tolong adakan rapat tentang sistem")
    mock_meeting = {
        "meeting_id": "test_123",
        "title": "Rapat",
        "participants": ["Alfa"],
        "dialogue_transcript": [],
        "execution_results": [],
    }
    with patch("alfa.bot.mode_commands.get_telegram_mode_sync", return_value="single"), \
         patch("alfa.swarm.engine.conduct_multi_agent_meeting", return_value=mock_meeting):
        res = conduct_ai_meeting(topic="Rapat kecanggihan sistem")
        assert res["status"] == "success"
        assert res["meeting_id"] == "test_123"


def test_tool_config_filters_conduct_meeting_in_single_mode():
    with patch("alfa.bot.mode_commands.get_telegram_mode_sync", return_value="single"), \
         patch("alfa.tools.rag.select_relevant_functions", side_effect=lambda tools, *args, **kwargs: tools):
        # In single mode without meeting keyword:
        tools1, _, _, _ = prepare_tools_and_configs("gemini-2.5", "sys", "halo apa kabar", [], False)
        tool_names1 = [getattr(t, "__name__", "") for t in tools1]
        assert "conduct_ai_meeting" not in tool_names1

        # With meeting keyword:
        tools2, _, _, _ = prepare_tools_and_configs("gemini-2.5", "sys", "adakan rapat tim", [], False)
        tool_names2 = [getattr(t, "__name__", "") for t in tools2]
        assert "conduct_ai_meeting" in tool_names2


def test_cancel_requested_auto_clears_when_no_session_running(tmp_path):
    flag_file = tmp_path / "swarm_cancel.flag"
    flag_file.write_text("cancel")
    with patch("alfa.swarm.workspace_hygiene.CANCEL_FLAG_FILE", str(flag_file)), \
         patch("alfa.swarm.workspace_hygiene._get_swarm_output_dir", return_value=str(tmp_path)), \
         patch("alfa.swarm.workspace_hygiene.any_session_running", return_value=False):
        workspace_hygiene.MEETING_RUNNING = False
        assert workspace_hygiene._cancel_requested() is False
        assert not flag_file.exists()


def test_wa_get_incoming_filters_broadcast_and_newsletter():
    mock_payload = {
        "status": "success",
        "count": 4,
        "incoming": [
            {"chat": "status@broadcast", "sender": "status@broadcast", "body": "status reaction"},
            {"chat": "12345@newsletter", "sender": "Channel News", "body": "breaking news"},
            {"chat": "628123456789@s.whatsapp.net", "sender": "Budi", "body": "Halo mas"},
            {"chat": "9876543210@g.us", "sender": "Andi", "body": "Rapat jam 2"},
        ]
    }
    with patch("alfa.tools.system.integrations._wa_api", return_value=mock_payload):
        res = wa_get_incoming(limit=10)
        assert res["status"] == "success"
        assert len(res["incoming"]) == 2
        assert res["incoming"][0]["sender"] == "Budi"
        assert res["incoming"][1]["sender"] == "Andi"
        assert res["count"] == 2


@pytest.mark.asyncio
async def test_wa_digest_formatting_preserves_markdown():
    mock_records = [
        {
            "format_name": "Laporan",
            "tab": "Laporan",
            "timestamp": "20/10/2026, 08.00.00",
            "sender": "Budi",
            "group": "Grup Operasional",
            "body": "Pekerjaan\tselesai\nAlat aman",
        }
    ]
    with patch("alfa.bot.wa_digest.collect_records", return_value=mock_records), \
         patch("alfa.bot.wa_digest.summarize_records", return_value="*Poin Utama:*\n- Item 1 sukses\n- Item 2 tertunda"):
        text = await build_digest_text()
        assert "*Poin Utama:*" in text
        assert "• **Laporan**" in text
        assert "Pekerjaan selesai Alat aman" in text


def test_get_active_api_key_sync_ignores_dummy_keys(isolated_db):
    from alfa.core.db import keys as k

    # Add dummy short key for custom provider
    k.add_api_key_sync("dummy", "custom_llm", "x", "m1")
    # Should not be returned as active valid key
    assert k.get_active_api_key_sync("custom_llm") is None

    # Add valid key
    k.add_api_key_sync("real", "custom_llm", "Secret123Valid", "m2")
    active = k.get_active_api_key_sync("custom_llm")
    assert active is not None
    assert active["api_key"] == "Secret123Valid"


@pytest.mark.asyncio
async def test_run_openai_agentic_turn_handles_missing_credentials_cleanly():
    from alfa.core.brain.agentic import run_openai_agentic_turn

    with patch("alfa.core.rate_limiter.check_rate_limit", return_value=(True, "")):
        # Missing api_key and empty base_url
        res = await run_openai_agentic_turn(
            provider="custom_unknown",
            base_url="",
            api_key="",
            model="custom-chat",
            system_instruction="sys",
            user_text="hello",
        )
        assert res is None


def test_get_agent_api_client_provides_default_base_url():
    from alfa.swarm.llm_client import get_agent_api_client

    agent = {"provider": "deepseek", "model": "deepseek-chat"}
    prov, key, model, base_url, key_id = get_agent_api_client(agent)
    assert prov == "deepseek"
    assert base_url == "https://api.deepseek.com/v1"


@pytest.mark.asyncio
async def test_multi_document_batching_executes_single_report():
    from alfa.bot.handlers import _execute_document_batch
    from unittest.mock import AsyncMock

    mock_docs = [
        {"file_name": "laporan_a.pdf", "mime_type": "application/pdf", "txt_ctx": "Data Omzet: Rp 50jt", "part": None},
        {"file_name": "laporan_b.pdf", "mime_type": "application/pdf", "txt_ctx": "Data Biaya: Rp 20jt", "part": None},
    ]
    batch = {
        "docs": mock_docs,
        "captions": ["minta rinciannya"],
        "user_id": 12345,
        "context": MagicMock(),
        "latest_update": MagicMock(),
        "turn_start_time": 0,
    }

    turn_prompts = []

    async def fake_run_turn(user_id, user_prompt, multimodal_parts, chat_id):
        turn_prompts.append(user_prompt)
        return "Laporan Rincian Terpadu: Laba bersih Rp 30jt."

    async def fake_typing(*a, **k):
        pass

    with patch("alfa.bot.handlers._run_agent_turn", side_effect=fake_run_turn), \
         patch("alfa.bot.handlers.safe_send_message", new_callable=AsyncMock) as mock_send, \
         patch("alfa.bot.handlers.send_typing_loop", side_effect=fake_typing):
        await _execute_document_batch(chat_id=999, batch=batch)

        # Harus tepat 1 panggilan ke agen dan 1 pesan balasan ke chat
        assert len(turn_prompts) == 1
        prompt = turn_prompts[0]
        assert "WAJIB JADIKAN 1 LAPORAN TERPADU DENGAN STRUKTUR RUNUT PER BERKAS" in prompt
        assert "FORMAT WAJIB: CUKUP NAMA DOKUMEN DAN RINCIAN MATERIALNYA SECARA RAPI" in prompt
        assert "laporan_a.pdf" in prompt
        assert "laporan_b.pdf" in prompt
        assert mock_send.call_count == 1
        sent_text = mock_send.call_args[0][2]
        assert "Laporan Rincian Terpadu" in sent_text



def test_permission_gate_dangerous_only_policy(monkeypatch):
    """
    Memverifikasi preferensi user: izin interaktif hanya diminta untuk perintah
    berisiko berbahaya (HIGH & CRITICAL). Tool dokumen, baca berkas, konversi,
    dan tool MEDIUM/LOW lainnya di-auto-approve tanpa merepotkan user satu per satu.
    """
    from alfa.core.perm import gate as g
    from alfa.core.perm.constants import RiskTier

    # 1. Saat PERMISSION_GATE_DANGEROUS_ONLY aktif (kondisi default di bot)
    monkeypatch.setattr(g, "PERMISSION_GATE_DANGEROUS_ONLY", True)

    # Trust score awal 0 / belum pernah ada persetujuan eksplisit
    monkeypatch.setattr(g, "get_tier_trust", lambda _cid, _tier: (0.0, 0))

    # Tool dokumen & pembacaan file (MEDIUM / LOW) harus otomatis lolos tanpa izin
    doc_tools = [
        "read_local_file",
        "analyze_dataset_csv_json",
        "markitdown_convert_document",
        "generate_excel_spreadsheet",
        "generate_pdf_report",
        "pdf_merge_documents",
        "libreoffice_convert_document",
        "convert_media_format",
        "web_search",
    ]
    for tool in doc_tools:
        ok, reason = g.should_auto_approve(8821693251, tool)
        assert ok is True, f"Tool {tool} seharusnya auto-approved"
        assert "auto_approved" in reason

    # Perintah berisiko berbahaya (HIGH) WAJIB tetap meminta izin
    high_tools = [
        "execute_bash_command",
        "execute_python_sandbox",
        "kill_process",
        "manage_system_services",
        "write_local_file",
        "query_database",
        "delete_dynamic_plugin",
    ]
    for tool in high_tools:
        ok, _ = g.should_auto_approve(8821693251, tool)
        assert ok is False, f"Perintah berisiko {tool} (HIGH) harus meminta izin"

    # Perintah irreversible & kritis (CRITICAL) WAJIB tetap meminta izin
    critical_tools = [
        "clean_system_storage",
        "ssh_execute_command",
        "control_linux_hardware",
        "manage_crontab_jobs",
        "vault_delete_secret",
    ]
    for tool in critical_tools:
        ok, _ = g.should_auto_approve(8821693251, tool)
        assert ok is False, f"Perintah berbahaya {tool} (CRITICAL) harus meminta izin"


def test_permission_gate_strict_mode_when_dangerous_only_false(monkeypatch):
    """
    Memverifikasi bila PERMISSION_GATE_DANGEROUS_ONLY dimatikan (mode ketat),
    tool MEDIUM kembali memerlukan akumulasi trust & approval eksplisit.
    """
    from alfa.core.perm import gate as g

    monkeypatch.setattr(g, "PERMISSION_GATE_DANGEROUS_ONLY", False)

    # Tanpa approval eksplisit -> MEDIUM ditolak
    monkeypatch.setattr(g, "get_tier_trust", lambda _cid, _tier: (1.0, 0))
    ok, _ = g.should_auto_approve(8821693251, "read_local_file")
    assert ok is False

    # Dengan approval eksplisit cukup (>=3) & trust >= 0.7 -> MEDIUM auto-approve
    monkeypatch.setattr(g, "get_tier_trust", lambda _cid, _tier: (0.8, 3))
    ok, _ = g.should_auto_approve(8821693251, "read_local_file")
    assert ok is True


def test_generate_excel_spreadsheet_dict_rows(tmp_path, monkeypatch):
    """Memverifikasi generate_excel_spreadsheet menerima rows_json bertipe list of dicts tanpa error openpyxl."""
    import openpyxl
    from alfa.tools.media.documents import generate_excel_spreadsheet
    import alfa.tools.media.documents as doc_mod

    monkeypatch.setattr(doc_mod, "SANDBOX_DIR", str(tmp_path))

    headers = ["No", "Cluster", "Nama Material", "QTY", "Satuan"]
    rows = [
        {"No": 1, "Cluster": "CLR56-99-263", "Nama Material": "Kabel ADSS 24c", "QTY": 3000, "Satuan": "Meter"},
        {"No": 2, "Cluster": "CLR56-99-263", "Nama Material": "16 Cores ODB", "QTY": 25, "Satuan": "Set"},
    ]
    res = generate_excel_spreadsheet(
        sheet_title="Rekap_Material",
        headers=headers,
        rows_json=json.dumps(rows),
        filename="test_rekap.xlsx",
    )
    assert res["status"] == "success"
    file_path = res["file_path"]
    assert os.path.exists(file_path)

    wb = openpyxl.load_workbook(file_path)
    ws = wb.active
    assert ws.max_row == 3  # Header + 2 rows
    assert ws.cell(row=2, column=2).value == "CLR56-99-263"
    assert ws.cell(row=2, column=3).value == "Kabel ADSS 24c"
    assert ws.cell(row=2, column=4).value == 3000


@pytest.mark.asyncio
async def test_handle_doc_action_callback_multi_doc_consolidated_prompt(monkeypatch):
    """Memverifikasi handle_doc_action_callback menghasilkan prompt konsolidasi 1 file Excel untuk batch multi-dokumen."""
    from alfa.bot.handlers import handle_doc_action_callback, _RECENT_DOCS

    chat_id = 998877
    _RECENT_DOCS[chat_id] = {
        "is_batch": True,
        "docs_count": 3,
        "file_name": "3 Berkas Gabungan (DO1.xls dkk)",
        "file_names": ["DO1.xls", "DO2.xls", "DO3.xls"],
        "mime_type": "multipart/batch",
        "txt_ctx": "Data DO1, DO2, DO3",
        "time": time.time(),
    }

    mock_update = MagicMock()
    mock_update.effective_chat.id = chat_id
    mock_update.effective_user.id = 12345
    mock_context = MagicMock()
    mock_context.bot.send_message = AsyncMock()

    captured_prompt = {}

    async def fake_run_turn(*args, **kwargs):
        captured_prompt["prompt"] = kwargs.get("user_prompt", "")
        return "Laporan Excel selesai dibuat."

    monkeypatch.setattr("alfa.bot.handlers._run_agent_turn", fake_run_turn)
    monkeypatch.setattr("alfa.bot.handlers.safe_send_message", AsyncMock())
    monkeypatch.setattr("alfa.bot.handlers._check_and_send_media_artifacts", AsyncMock())

    handled = await handle_doc_action_callback(mock_update, mock_context, "docact:excel")
    assert handled is True
    prompt = captured_prompt.get("prompt", "")
    assert "Rekap_Material_Gabungan.xlsx" in prompt
    assert "SATU (1) berkas spreadsheet Excel (.xlsx) gabungan" in prompt
    assert "DO1.xls, DO2.xls, DO3.xls" in prompt
    assert "DILARANG membuat file terpisah" in prompt


@pytest.mark.asyncio
async def test_check_and_send_media_artifacts_deduplicates_excel(tmp_path, monkeypatch):
    """Memverifikasi check_and_send_media_artifacts membuang data.xlsx jika ada file bernama spesifik dan hanya mengirim 1 file."""
    import alfa.bot.handlers as h_mod

    monkeypatch.setattr(h_mod, "SANDBOX_DIR", str(tmp_path))

    # Buat dua file Excel: data.xlsx (generic) dan Rekap_Material_Gabungan.xlsx (spesifik)
    data_file = tmp_path / "data.xlsx"
    data_file.write_text("dummy generic")
    rekap_file = tmp_path / "Rekap_Material_Gabungan.xlsx"
    rekap_file.write_text("dummy rekap")

    sent_docs = []

    mock_update = MagicMock()
    mock_update.effective_chat.id = 776655
    mock_context = MagicMock()

    async def fake_send_doc(chat_id, document, caption):
        sent_docs.append((caption, document.name))

    mock_context.bot.send_document = fake_send_doc

    await h_mod.check_and_send_media_artifacts(mock_update, mock_context)

    # Hanya Rekap_Material_Gabungan.xlsx yang dikirim!
    assert len(sent_docs) == 1
    assert "Rekap_Material_Gabungan.xlsx" in sent_docs[0][0]
    # data.xlsx harus terhapus dari disk
    assert not data_file.exists()


@pytest.mark.asyncio
async def test_swarm_cancel_stops_execution_immediately(tmp_path, monkeypatch):
    """Memverifikasi request_cancel_swarm menghentikan eksekusi agen seketika tanpa lanjut memproses tool."""
    from alfa.swarm import workspace_hygiene
    from alfa.swarm.meeting_steps import _execute_agent_step

    monkeypatch.setattr(workspace_hygiene, "_get_swarm_output_dir", lambda: str(tmp_path))

    sid = workspace_hygiene.register_session("swarm_test_cancel_123")
    assert workspace_hygiene.MEETING_RUNNING is True

    # Sebelum cancel, cancel_requested adalah False
    assert workspace_hygiene._cancel_requested(sid) is False

    # Panggil cancel
    ok = workspace_hygiene.request_cancel_swarm()
    assert ok is True
    assert workspace_hygiene._cancel_requested(sid) is True

    # Saat eksekusi agen dipanggil dalam kondisi cancel, harus langsung abort tanpa jalankan tool
    step_res, passed, feedback = await _execute_agent_step(
        {"name": "Code Crafter", "role": "Programmer"},
        "Buat skrip python",
        topic="Testing",
        intent_info={},
    )
    assert passed is False
    assert "DIBATALKAN" in step_res.get("execution_summary", "")

    # Cleanup
    workspace_hygiene.unregister_session(sid)
    workspace_hygiene._clear_cancel_flag(sid)


def test_format_swarm_result_when_cancelled():
    """Memverifikasi format_swarm_result menghasilkan teks pembatalan yang bersih tanpa laporan palsu sukses."""
    from alfa.bot.mode_commands import format_swarm_result

    steps_text, final_text, files = format_swarm_result({"status": "cancelled"})
    assert "dibatalkan" in steps_text.lower()
    assert "dihentikan" in final_text.lower()
    assert "🏆" not in final_text
    assert len(files) == 0




