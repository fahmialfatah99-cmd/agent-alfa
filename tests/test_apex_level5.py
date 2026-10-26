"""
Comprehensive Unit Tests for Level 5+ Sovereign Cognitive Apex Engine (ALFA OS).

Tests all 5 Apex Pillars:
1. Neuro-Symbolic Verification Engine (AST Math Prover & Invariant Solver)
2. Offline 'Dreaming Mode' & Deterministic Fast-Path Compiler
3. Predictive Operational World-Model (Pre-emptive Field Anomaly Detection)
4. Shadow-Git Canary Self-Evolution (3-Stage Gate & Automatic Canary Rollback)
5. Adaptive Hardware & Thermal Governor (Host Telemetry & Quarantine Enforcement)
"""

from alfa.core.apex import (
    analyze_field_event_preemptively,
    apex_neuro_symbolic_verify,
    apex_run_dreaming_consolidation,
    apex_synthesize_dynamic_tool,
    apex_system_telemetry,
    get_apex_full_status,
    get_canary_evolution_metrics,
    get_hardware_governor_state,
    get_world_model_summary,
    is_service_quarantined,
    run_idle_dreaming_cycle,
    safe_ast_math_eval,
    synthesize_and_hotload_tool,
    try_execute_compiled_fastpath,
    verify_and_prove_response,
    verify_arithmetic_equations_in_text,
    verify_code_via_canary_gate,
)


def test_pillar1_neuro_symbolic_ast_prover():
    # Safe AST evaluation
    assert safe_ast_math_eval("68 - 7") == 61
    assert safe_ast_math_eval("63 + 94 + 1 + 2") == 160
    assert safe_ast_math_eval("(68 - 7) + 99") == 160
    # Reject unsafe non-math strings
    assert safe_ast_math_eval("__import__('os').system('ls')") is None

    # Auto-repair miscalculated inline equations
    raw_text = "Mandor Mugi mengambil 68 - 7 = 64 batang dan total 61 + 99 = 170 batang."
    fixed_text, findings = verify_arithmetic_equations_in_text(raw_text)
    assert "68 - 7 = 61" in fixed_text
    assert "61 + 99 = 160" in fixed_text
    assert len(findings) == 2
    assert not findings[0]["is_valid"]
    assert not findings[1]["is_valid"]

    # Domain Invariant Prover for Lahat
    res = verify_and_prove_response(
        user_prompt="Berapa total tiang Lahat?",
        reply_text="Total bersih adalah 145 batang untuk cluster Lahat.",
    )
    assert not res["is_valid"]
    assert "160 batang" in res["corrected_reply"]
    assert len(res["symbolic_proofs"]) >= 1


def test_pillar2_dreaming_fastpath_compiler():
    cycle_res = run_idle_dreaming_cycle(force=True)
    assert cycle_res["status"] == "success"
    assert cycle_res["seeded_fastpaths"] >= 1

    # Sub-millisecond fastpath execution for Lahat pole summary
    fp_hit = try_execute_compiled_fastpath("berapa total tiang bersih lahat?")
    assert fp_hit is not None
    assert "160 Batang" in fp_hit["response"]
    assert fp_hit["saved_tokens"] >= 800

    # Never trigger fastpath when user is issuing a correction
    fp_corr = try_execute_compiled_fastpath("salah, bukan begitu total tiang lahat")
    assert fp_corr is None


def test_pillar3_predictive_world_model():
    # Detect unquantified return risk
    alert1 = analyze_field_event_preemptively("Mandor Mugi ada retur tiang hari ini di Lahat")
    assert alert1["status"] == "alert"
    assert any(a["category"] == "UNQUANTIFIED_RETURN" for a in alert1["anomalies_detected"])

    # Detect large pole dispatch without breakdown specification
    alert2 = analyze_field_event_preemptively("Pengambilan 85 batang tiang untuk cluster Curup")
    assert alert2["status"] == "alert"
    assert any(a["category"] == "MISSING_POLE_SPEC_BREAKDOWN" for a in alert2["anomalies_detected"])

    # Clean dispatch with full specification passes without alert
    clean_ev = analyze_field_event_preemptively("Pengambilan 60 batang tiang T73 7m 3 inch Lahat")
    assert clean_ev["status"] == "clean"

    summary = get_world_model_summary()
    assert summary["status"] == "success"
    assert summary["unresolved_anomalies_count"] >= 2


def test_pillar4_canary_evolution_and_rollback():
    # 1. Reject code touching quarantined service wa-sheets-bot
    bad_code = "def bad_tool():\n    return 'systemctl restart wa-sheets-bot.service'\n"
    gate_bad = verify_code_via_canary_gate("bad_tool", bad_code)
    assert not gate_bad["passed"]
    assert "QUARANTINE" in gate_bad["stage"]

    # 2. Synthesize & hotload valid dynamic tool with smoke test
    valid_code = (
        "def calc_pole_reserve(net_poles: int, buffer_pct: float = 0.1) -> dict:\n"
        "    reserve = int(round(net_poles * buffer_pct))\n"
        "    return {'net': net_poles, 'reserve': reserve, 'total_with_buffer': net_poles + reserve}\n"
    )
    syn_res = synthesize_and_hotload_tool(
        tool_name="calc_pole_reserve",
        description="Calculate pole reserve buffer for FTTH cluster",
        code_source=valid_code,
        smoke_test_args={"net_poles": 160, "buffer_pct": 0.1},
    )
    assert syn_res["status"] == "success"
    assert syn_res["smoke_output"]["total_with_buffer"] == 176

    # 3. Verify Canary Auto-Rollback when a dynamic tool fails at runtime
    faulty_runtime_code = (
        "def unstable_sensor_tool(x: int) -> int:\n"
        "    if x < 0:\n"
        "        raise RuntimeError('Sensor calibration underflow')\n"
        "    return x * 2\n"
    )
    gate_ok = verify_code_via_canary_gate(
        "unstable_sensor_tool",
        faulty_runtime_code,
        smoke_test_args={"x": 5},
    )
    assert gate_ok["passed"]
    synthesize_and_hotload_tool(
        tool_name="unstable_sensor_tool",
        description="Test canary rollback on negative input",
        code_source=faulty_runtime_code,
        smoke_test_args={"x": 5},
    )

    from alfa.core.apex.canary_evolution import _wrap_with_canary_rollback_guard
    guarded = _wrap_with_canary_rollback_guard("unstable_sensor_tool", gate_ok["callable"])
    rb_res = guarded(x=-1)
    assert rb_res["status"] == "canary_rolled_back"
    assert "Sensor calibration underflow" in rb_res["error"]

    metrics = get_canary_evolution_metrics()
    assert metrics["total_canary_rollbacks"] >= 1


def test_pillar5_hardware_governor_and_agent_tools():
    gov = get_hardware_governor_state()
    assert gov["status"] == "success"
    assert gov["governor_mode"] in ("APEX_TURBO", "BALANCED", "THERMAL_PROTECTION")
    assert gov["cpu_temp_celsius"] > 0

    # Verify strict quarantine on wa-sheets-bot.service
    assert is_service_quarantined("wa-sheets-bot.service")
    assert is_service_quarantined("wa-sheets-bot")
    assert not is_service_quarantined("telegram-ai-bot.service")

    # Test registered agent tools
    ns_tool = apex_neuro_symbolic_verify("68 - 7 = 61", "Lahat")
    assert ns_tool["status"] == "success"

    dr_tool = apex_run_dreaming_consolidation()
    assert dr_tool["status"] == "success"

    tel_tool = apex_system_telemetry()
    assert tel_tool["status"] == "success"
    assert "Level 5+" in tel_tool["autonomy_level"]

    syn_tool = apex_synthesize_dynamic_tool(
        tool_name="km_to_drum_estimator",
        description="Estimate cable drums needed for route km",
        python_code="def km_to_drum_estimator(km: float) -> int:\n    return int((km * 1000) // 4000) + 1\n",
    )
    assert syn_tool["status"] == "success"

    full = get_apex_full_status()
    assert full["status"] == "success"
    assert "cloud_rescue_vault" in full
    assert "symbolic_ledgers" in full


def test_perfection_4_phases_vault_ledger_and_maintenance(tmp_path):
    import openpyxl
    from alfa.core.apex import (
        audit_excel_formula_integrity,
        run_autonomous_apex_maintenance,
        stage_file_for_cloud_rescue,
    )

    # Phase 1: Cloud-Local Rescue Vault
    dummy_file = tmp_path / "Master_Laporan_Cluster_Lahat_Test.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rekap"
    ws["G6"] = 61
    ws["G7"] = 99
    ws["G8"] = "=SUM(G6:G7)"  # Clean dynamic formula (no circular ref)
    wb.save(str(dummy_file))

    vault_res = stage_file_for_cloud_rescue(
        str(dummy_file),
        cluster_name="LAHAT",
        attempt_cloud_upload=False,
        vault_root=str(tmp_path / "vault"),
    )
    assert vault_res["status"] == "success"
    assert vault_res["sync_status"] == "LOCAL_VAULTED"

    # Phase 2: Symbolic Excel Formula Integrity Auditor (Clean vs Circular)
    audit_ok = audit_excel_formula_integrity(str(dummy_file))
    assert audit_ok["status"] == "valid"
    assert audit_ok["is_clean"] is True
    assert audit_ok["formula_count"] == 1

    # Test Circular Reference Detection
    bad_xlsx = tmp_path / "Bad_Circular.xlsx"
    wb_bad = openpyxl.Workbook()
    ws_bad = wb_bad.active
    ws_bad["G10"] = "=SUM(G6:G10)"  # Circular self-reference!
    wb_bad.save(str(bad_xlsx))

    audit_bad = audit_excel_formula_integrity(str(bad_xlsx))
    assert audit_bad["status"] == "circular_detected"
    assert audit_bad["is_clean"] is False
    assert len(audit_bad["circular_references"]) == 1

    # Phase 3 & 4: Full Autonomous Maintenance Cycle
    maint = run_autonomous_apex_maintenance(force=True)
    assert maint["status"] == "success"
    assert maint["sqlite_wal_checkpoint"] == "ok"


def test_universal_computer_use_engine():
    from alfa.core.apex import (
        apex_computer_use_action,
        execute_gui_action,
        inspect_desktop_environment,
    )

    env_info = inspect_desktop_environment()
    assert env_info["status"] == "success"
    assert "screen_resolution" in env_info

    # Vision Grid Screenshot capture
    cap_res = apex_computer_use_action(action="capture_grid")
    assert cap_res["status"] == "success"
    assert cap_res["with_grid"] is True

    # Simulated click action
    click_res = execute_gui_action(action="click", x=400, y=300, dry_run=True)
    assert click_res["status"] == "simulated"
    assert click_res["coordinates"] == {"x": 400, "y": 300}

    # Quarantine guard blocks interacting with wa-sheets-bot
    blocked_res = execute_gui_action(action="type", text="kill wa-sheets-bot", dry_run=True)
    assert blocked_res["status"] == "blocked_by_quarantine"


def test_temporal_graphrag_multihop_second_brain():
    from alfa.core.apex import (
        apex_temporal_graphrag_query,
        ingest_text_to_temporal_graph,
    )

    # Link A -> B and B -> C to test 2-Hop Bridge Discovery
    apex_temporal_graphrag_query(
        "Project Orion uses FastAPI Backend",
        mode="link",
        relation_source="PROJECT_ORION",
        relation_type="USES",
        relation_target="FASTAPI_BACKEND",
    )
    apex_temporal_graphrag_query(
        "FastAPI Backend connects to SQLite WAL Engine",
        mode="link",
        relation_source="FASTAPI_BACKEND",
        relation_type="CONNECTED_TO",
        relation_target="SQLITE_WAL_ENGINE",
    )

    # Auto-ingest natural text
    ing_res = ingest_text_to_temporal_graph("Sistem QuantumMesh menggunakan RustEngine untuk komputasi cepat.")
    assert ing_res["status"] == "success"
    assert ing_res["relations_processed"] >= 1

    # Query seed 'PROJECT_ORION': should discover 1-hop 'FASTAPI_BACKEND' and 2-hop 'SQLITE_WAL_ENGINE'
    q_res = apex_temporal_graphrag_query("PROJECT_ORION", mode="query")
    assert q_res["status"] == "success"
    assert any(r["target_entity"] == "FASTAPI_BACKEND" for r in q_res["hop1_relations"])
    assert any(r["target_entity"] == "SQLITE_WAL_ENGINE" for r in q_res["hop2_bridges"])


def test_autonomous_software_factory_and_e2e_verifier(tmp_path):
    from alfa.core.apex import build_and_verify_fullstack_app

    # Provide partial HTML (missing <!DOCTYPE> and <html> wrapper) to test Self-Healing + 4-Stage E2E Gate
    raw_html = """
    <div id="app">
        <h1>Autonomous SaaS Analytics</h1>
        <form id="query-form">
            <input type="text" name="metric" placeholder="Metric name" />
            <button type="submit">Run Analysis</button>
        </form>
    </div>
    <script>
        console.log("Analytics Ready");
    </script>
    """
    backend_code = """
def get_analytics_summary() -> dict:
    return {"status": "ok", "active_users": 42}
"""

    build_res = build_and_verify_fullstack_app(
        app_slug="saas-analytics-pro",
        title="SaaS Analytics Pro",
        frontend_html=raw_html,
        backend_py=backend_code,
        auto_heal=True,
        sites_root=str(tmp_path / "sites"),
    )
    assert build_res["status"] == "success"
    assert build_res["e2e_verification"]["status"] == "PRODUCTION_READY"
    assert build_res["e2e_verification"]["e2e_score"] == 100
    assert build_res["e2e_verification"]["live_http_smoke"]["passed"] is True
    assert build_res["e2e_verification"]["live_http_smoke"]["http_status"] == 200
    assert len(build_res["e2e_verification"]["auto_repairs_applied"]) >= 1


def test_audit_fixes_tool_rag_sites_graphrag_and_rea_clone(tmp_path):
    from alfa.core.apex import (
        format_temporal_graphrag_prompt,
        rea_clone_and_build_app,
    )
    from alfa.tools import AVAILABLE_TOOLS
    from alfa.tools.rag import select_relevant_functions

    # 1. Verify Celah 1 Fix: Tool-RAG selects REA + Software Factory tools within TOP_K=24
    selected_fns = select_relevant_functions(
        AVAILABLE_TOOLS,
        "Tolong bedah aplikasi ini dengan rea dan buatkan aplikasi web sendiri di /sites",
        top_k=24,
    )
    selected_names = {getattr(fn, "__name__", "") for fn in selected_fns}
    assert "rea_clone_and_build_app" in selected_names
    assert "apex_software_factory_build" in selected_names
    assert "rea_synthesize_skill_from_code" in selected_names
    assert "serve_site" in selected_names

    # 2. Verify Celah 3 Fix: Automatic 2-Hop Temporal GraphRAG prompt formatting
    graph_prompt = format_temporal_graphrag_prompt("Ceritakan tentang PROJECT_ORION dan arsitekturnya")
    assert "TEMPORAL GRAPHRAG" in graph_prompt
    assert "FASTAPI_BACKEND" in graph_prompt
    assert "SQLITE_WAL_ENGINE" in graph_prompt

    # 3. Verify Celah 2 & Celah 4 Fix: 1-Click REA Dissect -> Learn -> Build -> Live HTTP 200 Test
    sample_js = tmp_path / "target_bundle.js"
    sample_js.write_text(
        "const API_URL = 'https://api.acme-logistics.io/v2/dispatch';\n"
        "function calculateTariff(weightKg) { return weightKg * 14500; }\n",
        encoding="utf-8",
    )
    clone_res = rea_clone_and_build_app(
        target_file_path=str(sample_js),
        app_slug="acme-logistics-clone",
        app_title="ACME Logistics Clone",
    )
    assert clone_res["status"] == "success"
    assert clone_res["rea_analysis"]["status"] == "success"
    assert clone_res["build_result"]["e2e_verification"]["live_http_smoke"]["http_status"] == 200
    assert clone_res["build_result"]["registered_in_sites_menu"] is True
    assert "/serve acme-logistics-clone" in clone_res["serve_command"]


def test_devin_class_giant_repo_engine_and_atomic_rollback(tmp_path):
    from alfa.core.apex import (
        build_and_verify_multifile_repo,
        parse_stack_trace_locations,
    )
    from alfa.core.cli.agent_engine.repomap import RepomapGenerator
    from alfa.tools import AVAILABLE_TOOLS
    from alfa.tools.filesystem.lsp_code_intelligence import (
        lsp_analyze_module_hierarchy,
        lsp_find_symbol_definition,
    )
    from alfa.tools.rag import select_relevant_functions

    # 1. Multi-File Repository Build + Automated Unittest Execution + Live HTTP 200 OK
    manifest_v1 = {
        "src/__init__.py": "",
        "src/core/__init__.py": "",
        "src/core/calculator.py": (
            "def compute_margin(revenue: float, cost: float) -> float:\n"
            "    return round(revenue - cost, 2)\n"
        ),
        "src/api/controller.py": (
            "from src.core.calculator import compute_margin\n\n"
            "def handle_request(rev: float, cost: float) -> dict:\n"
            "    return {'margin': compute_margin(rev, cost)}\n"
        ),
        "frontend/widget.ts": (
            "import { api } from './controller';\n"
            "export interface MarginResult { margin: number; }\n"
            "export function formatCurrency(val: number): string { return `Rp ${val}`; }\n"
        ),
        "tests/test_calculator.py": (
            "import unittest\n"
            "from src.api.controller import handle_request\n\n"
            "class TestCalc(unittest.TestCase):\n"
            "    def test_margin(self):\n"
            "        self.assertEqual(handle_request(100.0, 35.5)['margin'], 64.5)\n"
        ),
        "index.html": "<div><h1>Enterprise Repo Dashboard</h1><button>Test</button></div>",
    }

    res_ok = build_and_verify_multifile_repo(
        repo_slug="enterprise-erp-core",
        title="Enterprise ERP Core",
        files_manifest=manifest_v1,
        test_command="python3 -m unittest discover -s tests",
        atomic_rollback_on_failure=True,
        sites_root=str(tmp_path / "repos"),
    )
    assert res_ok["status"] == "success"
    assert res_ok["e2e_verification"]["e2e_score"] == 100
    assert res_ok["e2e_verification"]["atomic_rolled_back"] is False
    assert res_ok["e2e_verification"]["live_http_smoke"]["http_status"] == 200

    repo_dir = res_ok["repo_dir"]

    # 2. Polyglot RepoMap Centrality & LSP on TypeScript + Python
    gen = RepomapGenerator(root_dir=tmp_path / "repos" / "enterprise-erp-core")
    ranked = gen.rank_files_by_centrality()
    assert len(ranked) >= 5
    # calculator.py is imported by controller.py, so its inbound_refs >= 1
    calc_entry = next(r for r in ranked if str(r["rel_file"]).endswith("calculator.py"))
    assert int(calc_entry["inbound_refs"]) >= 1

    ts_sym = lsp_find_symbol_definition("formatCurrency", repo_path=repo_dir)
    assert ts_sym["status"] == "success"
    assert ts_sym["total_definitions"] == 1

    ts_hier = lsp_analyze_module_hierarchy("frontend/widget.ts", repo_path=repo_dir)
    assert ts_hier["status"] == "success"
    assert ts_hier["total_classes"] >= 1

    # 3. Atomic Rollback Verification on Bad Multi-File Patch
    bad_patch = {
        "src/core/calculator.py": "def compute_margin(revenue, cost)\n    return broken_syntax!!!\n"
    }
    res_bad = build_and_verify_multifile_repo(
        repo_slug="enterprise-erp-core",
        title="Enterprise ERP Core",
        files_manifest=bad_patch,
        atomic_rollback_on_failure=True,
        sites_root=str(tmp_path / "repos"),
    )
    assert res_bad["status"] == "rolled_back"
    assert res_bad["e2e_verification"]["atomic_rolled_back"] is True
    assert any(loc["file"].endswith("calculator.py") for loc in res_bad["e2e_verification"]["stack_trace_locations"])

    # Verify original file on disk is 100% intact after rollback!
    with open(f"{repo_dir}/src/core/calculator.py", encoding="utf-8") as f:
        restored_code = f.read()
    assert "round(revenue - cost, 2)" in restored_code

    # 4. Stack trace parser check
    parsed_locs = parse_stack_trace_locations('File "src/api/controller.py", line 14, in handle_request\nfrontend/widget.ts:29:5: error TS2322')
    assert len(parsed_locs) == 2

    # 5. Tool-RAG verification for Giant Repo tools
    selected = select_relevant_functions(
        AVAILABLE_TOOLS,
        "Tolong analisis arsitektur repositori ini dengan repomap dan lsp lalu refactor kode di sandbox",
        top_k=24,
    )
    sel_names = {getattr(fn, "__name__", "") for fn in selected}
    assert "generate_codebase_repomap" in sel_names
    assert "lsp_find_symbol_definition" in sel_names
    assert "git_worktree_sandbox_create" in sel_names
    assert "apex_repo_architect_build" in sel_names


def test_deep_audit_5_fixes_permissions_swarm_and_lowercase_graphrag():
    from alfa.bot.pipeline.tool_config import prepare_tools_and_configs
    from alfa.core.apex import ingest_text_to_temporal_graph
    from alfa.core.perm.constants import SAFE_TOOLS, TOOL_CLASSIFICATION
    from alfa.swarm.llm_client import SWARM_HEADLESS_ALLOWED, get_agent_api_client
    from alfa.tools import AVAILABLE_TOOLS

    # 1. Zero unclassified tools in AVAILABLE_TOOLS
    unclassified = [
        getattr(f, "__name__", "")
        for f in AVAILABLE_TOOLS
        if getattr(f, "__name__", "") not in TOOL_CLASSIFICATION
        and getattr(f, "__name__", "") not in SAFE_TOOLS
    ]
    assert unclassified == [], f"Unclassified tools found: {unclassified}"

    # 2. Swarm Headless Gate includes Software Factory & REA tools
    assert "apex_software_factory_build" in SWARM_HEADLESS_ALLOWED
    assert "apex_repo_architect_build" in SWARM_HEADLESS_ALLOWED
    assert "rea_clone_and_build_app" in SWARM_HEADLESS_ALLOWED

    # 3. Swarm Agent API Client automatic fallback when provider has no key
    dummy_agent = {"id": 999, "name": "Code Crafter", "provider": "nonexistent_provider", "model": "coder-x"}
    prov, api_key, model, _url, _kid = get_agent_api_client(dummy_agent)
    assert bool(api_key), "Swarm agent should automatically fall back to an active vault key"

    # 4. Lowercase + Markdown Temporal GraphRAG extraction
    ing = ingest_text_to_temporal_graph(
        "tolong cek pengambilan material di grup yofc tjhin emr bengkulu dan aplikasi *Concat Pro Studio*"
    )
    assert ing["status"] == "success"
    assert ing["entities_processed"] >= 2

    # 5. Swarm & Subagent unlocked in tool_config when user asks to build/clone an app
    tools_selected, _cfg, _cfg_f, _cfg_w = prepare_tools_and_configs(
        model_name="gemini-3.8-flash",
        system_instruction="sys",
        user_prompt="buat aplikasi web dan clone dengan tim agen",
        history_msgs=[],
        action_intent=True,
    )
    tool_names = {getattr(t, "__name__", "") for t in tools_selected}
    assert "apex_software_factory_build" in tool_names


def test_swarm_adaptive_dual_persona_and_smart_routing():
    from unittest.mock import patch

    from alfa.bot.mode_commands import format_swarm_result
    from alfa.core.runtime_ctx import current_user_prompt_var
    from alfa.swarm.live_logger import detect_task_intent
    from alfa.swarm.personas import AGENTS, get_adaptive_system_instruction
    from alfa.tools.system.management import conduct_ai_meeting

    # 1. All 7 Swarm Agents defined in AGENTS have Dual-Persona instructions
    for aid in (1, 2, 3, 4, 5, 6, 9):
        assert aid in AGENTS
        assert "Persona Umum" in AGENTS[aid]["system_instruction"]
        assert "Persona Engineering" in AGENTS[aid]["system_instruction"]

    # 2. General problem intent detection & Mode 2 persona activation
    gen_topic = "Rapatkan cara mengatasi keterlambatan pengiriman material kabel dan mandor di proyek Bengkulu"
    gen_intent = detect_task_intent(gen_topic)
    assert gen_intent["is_general_advisory"] is True
    assert gen_intent["is_build"] is False
    assert gen_intent["recommended_mode"] == "plan"

    code_crafter_gen_sys = get_adaptive_system_instruction(
        {"id": 2, "name": "Code Crafter"},
        topic=gen_topic,
        mode="plan",
        intent_info=gen_intent,
    )
    assert "MODE 2 — GENERAL PROBLEM-SOLVING" in code_crafter_gen_sys
    assert "DILARANG KERAS membuat website" in code_crafter_gen_sys

    # 3. Coding/Build intent detection & Mode 1 persona activation
    build_topic = "Tim swarm buatkan aplikasi website manajemen stok gudang dengan FastAPI"
    build_intent = detect_task_intent(build_topic)
    assert build_intent["is_build"] is True
    assert build_intent["is_general_advisory"] is False
    assert build_intent["recommended_mode"] == "execute"

    code_crafter_build_sys = get_adaptive_system_instruction(
        {"id": 2, "name": "Code Crafter"},
        topic=build_topic,
        mode="execute",
        intent_info=build_intent,
    )
    assert "MODE 1 — ENGINEERING, BUILD" in code_crafter_build_sys

    # 4. conduct_ai_meeting smart routing: general problem -> mode='plan', build -> mode='execute'
    current_user_prompt_var.set("adakan rapat tim untuk masalah operasional")
    captured_modes = []

    async def _fake_conduct(topic, part_list=None, rounds=2, mode="execute", target_folder=""):
        captured_modes.append((topic, mode, rounds))
        return {
            "meeting_id": "m_1",
            "title": "Rapat",
            "participants": ["Alpha Lead", "Code Crafter"],
            "dialogue_transcript": [
                {
                    "agent_name": "Alpha Lead",
                    "role": "Moderator",
                    "avatar_emoji": "👑",
                    "message": "Solusi utama adalah audit jadwal pengiriman.",
                }
            ],
            "execution_results": [],
            "consensus": "Perketat SOP pengiriman material.",
            "action_plan": "1. Cek jadwal mandor.",
        }

    with patch("alfa.bot.mode_commands.get_telegram_mode_sync", return_value="single"), \
         patch("alfa.swarm.engine.conduct_multi_agent_meeting", side_effect=_fake_conduct):
        res_gen = conduct_ai_meeting(topic=gen_topic)
        assert res_gen["mode"] == "plan"
        assert captured_modes[-1][1] == "plan"
        assert captured_modes[-1][2] >= 2

        res_build = conduct_ai_meeting(topic=build_topic)
        assert res_build["mode"] == "execute"
        assert captured_modes[-1][1] == "execute"

    # 5. format_swarm_result renders general problem discussion & action plan cleanly
    s_txt, f_txt, files = format_swarm_result(res_gen)
    assert "Ringkasan Diskusi" in s_txt
    assert "KONSENSUS & SOLUSI TIM SWARM" in f_txt
    assert "ACTION PLAN" in f_txt
    assert files == []


def test_fact_grounded_war_room_meeting_export_and_db_snapshot_manager(tmp_path):
    import os

    from alfa.core.apex import ingest_text_to_temporal_graph
    from alfa.core.db.backup_manager import (
        create_verified_db_snapshot,
        list_db_snapshots,
        restore_db_snapshot,
        verify_sqlite_integrity,
    )
    from alfa.swarm.deliberation import (
        export_meeting_report_files,
        gather_pre_meeting_intel,
        load_last_meeting_state,
    )

    # 1. Fact-Grounded War Room pre-meeting intel from GraphRAG
    ingest_text_to_temporal_graph(
        "proyek *Cluster Lahat* dikelola oleh mandor *Aqe Setiawan* untuk pemasangan *Tiang 7m*"
    )
    intel = gather_pre_meeting_intel(
        "Rapatkan evaluasi pemasangan Tiang 7m di Cluster Lahat oleh mandor Aqe Setiawan",
        allow_web_search=False,
    )
    assert "LEMBAR FAKTA PRA-RAPAT" in intel
    assert "LAHAT" in intel.upper() or "TIANG" in intel.upper()

    # 2. Meeting Report Exporter (.md & .html) + Last Meeting State Persistence
    meetings_dir = str(tmp_path / "meetings")
    sample_meeting = {
        "session_id": "swarm_test_99",
        "topic": "Strategi Efisiensi Pengiriman Material Proyek Lahat",
        "mode": "plan",
        "participants": ["Alpha Lead", "System Auditor", "Strategic Planner"],
        "pre_meeting_intel": intel,
        "consensus": "Gunakan jadwal pengiriman terpusat setiap Senin pagi.",
        "action_plan": "1. Verifikasi stok gudang.\n2. Konfirmasi mandor lapangan.",
        "vote_tally": {
            "winner": {
                "id": "p1-strategic-planner",
                "author": "Strategic Planner",
                "votes_for": 3,
                "votes_against": 0,
                "summary": "Jadwal terpusat Senin pagi",
            },
            "agreement": 1.0,
        },
        "dialogue_transcript": [
            {
                "round": 1,
                "agent_name": "Strategic Planner",
                "role": "Strategist",
                "avatar_emoji": "💡",
                "message": "Pengiriman terpusat menghemat biaya logistik 25%.",
            }
        ],
    }
    exported = export_meeting_report_files(sample_meeting, out_dir=meetings_dir)
    assert os.path.exists(exported["markdown"])
    assert os.path.exists(exported["html"])
    with open(exported["html"], encoding="utf-8") as hf:
        html_txt = hf.read()
    assert "DOKUMEN KEPUTUSAN RAPAT AI SWARM" in html_txt
    assert "Gunakan jadwal pengiriman terpusat" in html_txt

    loaded_state = load_last_meeting_state(out_dir=meetings_dir)
    assert loaded_state["session_id"] == "swarm_test_99"
    assert "Verifikasi stok gudang" in loaded_state["action_plan"]

    # 3. Verified SQLite Snapshot & Auto-Rotation Manager
    bkp_dir = str(tmp_path / "backups")
    health = verify_sqlite_integrity()
    assert health["healthy"] is True

    snap1 = create_verified_db_snapshot(backup_dir=bkp_dir, max_backups=2)
    assert snap1["status"] == "success"
    assert os.path.exists(snap1["snapshot_path"])

    snap2 = create_verified_db_snapshot(backup_dir=bkp_dir, max_backups=2)
    snap3 = create_verified_db_snapshot(backup_dir=bkp_dir, max_backups=2)
    assert snap2["status"] == "success"
    assert snap3["status"] == "success"

    listed = list_db_snapshots(backup_dir=bkp_dir)
    assert listed["total"] == 2  # Rotated automatically to max_backups=2

    restored_db_path = str(tmp_path / "restored_agent_data.db")
    rest = restore_db_snapshot(snap3["snapshot_path"], target_db_path=restored_db_path)
    assert rest["status"] == "success"
    assert os.path.exists(restored_db_path)
    assert verify_sqlite_integrity(restored_db_path)["healthy"] is True







