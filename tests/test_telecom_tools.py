"""
Unit tests for Telecommunication & Fiber Optic (NRO / FTTH) Operations Suite.
"""

import os

import openpyxl

from alfa.tools.telecom_tools import (
    _classify_telecom_material,
    boq_check_project_anomalies,
    boq_compare_and_diff,
    boq_consolidate_master_workbook,
    boq_reconcile_with_field_reports,
    render_cluster_kml_map,
)


def test_classify_telecom_material():
    assert _classify_telecom_material("Kabel ADSS 24 Core") == "CABLE"
    assert _classify_telecom_material("Steel Pole 7 meter 3 inch") == "POLE"
    assert _classify_telecom_material("FAT 16 Port Pole Mounted") == "DISTRIBUTION_BOX"
    assert _classify_telecom_material("Suspension Clamp Plastic") == "ACCESSORIES"
    assert _classify_telecom_material("Protection Sleeve 60mm") == "OPTICAL_PASSIVE"
    assert _classify_telecom_material("Random Item ABC") == "OTHER"


def test_render_cluster_kml_map(tmp_path):
    kml_content = """<?xml version="1.0" encoding="UTF-8"?>
    <kml xmlns="http://www.opengis.net/kml/2.2">
      <Document>
        <Placemark>
          <name>Pole 7m 01</name>
          <description>Tiang awal cluster</description>
          <Point>
            <coordinates>102.54914,-3.45673,0</coordinates>
          </Point>
        </Placemark>
        <Placemark>
          <name>FAT 01</name>
          <description>FAT 16 port</description>
          <Point>
            <coordinates>102.54950,-3.45690,0</coordinates>
          </Point>
        </Placemark>
        <Placemark>
          <name>Rute Kabel Feeder</name>
          <LineString>
            <coordinates>
              102.54914,-3.45673,0 102.54950,-3.45690,0
            </coordinates>
          </LineString>
        </Placemark>
      </Document>
    </kml>
    """
    kml_file = tmp_path / "test_cluster.kml"
    kml_file.write_text(kml_content, encoding="utf-8")

    out_html = "test_cluster_map.html"
    res = render_cluster_kml_map(str(kml_file), output_html_filename=out_html)

    assert res["status"] == "success"
    assert res["total_points"] == 2
    assert res["total_cable_routes"] == 1
    assert os.path.exists(res["output_html_file"])


def test_boq_consolidate_master_workbook(tmp_path):
    # Buat Workbook 1 (Cluster A)
    wb1 = openpyxl.Workbook()
    ws1 = wb1.active
    ws1.title = "BoQ NRO Cluster"
    ws1.append(["No", "Description", "Unit", "Qty"])
    ws1.append([1, "Kabel ADSS 24C", "Meter", 1200])
    ws1.append([2, "Steel Pole 7m", "Pcs", 20])
    wb1_path = tmp_path / "BOQ - Cluster A.xlsx"
    wb1.save(str(wb1_path))

    # Buat Workbook 2 (Cluster B)
    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2.title = "BoQ NRO Cluster"
    ws2.append(["No", "Description", "Unit", "Qty"])
    ws2.append([1, "Kabel ADSS 24C", "Meter", 800])
    ws2.append([2, "Steel Pole 7m", "Pcs", 15])
    ws2.append([3, "FAT 16 Core", "Box", 4])
    wb2_path = tmp_path / "BOQ - Cluster B.xlsx"
    wb2.save(str(wb2_path))

    res = boq_consolidate_master_workbook(
        file_paths=[str(wb1_path), str(wb2_path)],
        output_filename="TEST_MASTER_CONSOLIDATED.xlsx",
        project_title="TEST PROJECT KONSOLIDASI",
    )

    assert res["status"] == "success"
    assert res["total_clusters"] == 2
    assert res["grand_totals"]["cable_meters"] == 2000.0
    assert res["grand_totals"]["poles_pcs"] == 35
    assert res["grand_totals"]["fat_fdt_pcs"] == 4
    assert os.path.exists(res["output_file"])


def test_boq_compare_and_diff(tmp_path):
    wb_old = openpyxl.Workbook()
    ws_old = wb_old.active
    ws_old.append(["Item", "Qty", "Unit"])
    ws_old.append(["Kabel ADSS 24C", 1000, "Meter"])
    ws_old.append(["Tiang 7m", 25, "Pcs"])
    ws_old.append(["FAT 8 Core", 5, "Pcs"])
    old_path = tmp_path / "BOQ_Rev0.xlsx"
    wb_old.save(str(old_path))

    wb_new = openpyxl.Workbook()
    ws_new = wb_new.active
    ws_new.append(["Item", "Qty", "Unit"])
    ws_new.append(["Kabel ADSS 24C", 1250, "Meter"])  # Modified (+250)
    ws_new.append(["Tiang 7m", 25, "Pcs"])            # Unchanged
    # FAT 8 Core removed
    ws_new.append(["FAT 16 Core", 3, "Pcs"])          # Added
    new_path = tmp_path / "BOQ_Rev1.xlsx"
    wb_new.save(str(new_path))

    res = boq_compare_and_diff(str(old_path), str(new_path), output_diff_filename="TEST_DIFF.xlsx")

    assert res["status"] == "success"
    stats = res["statistics"]
    assert stats["items_added"] == 1
    assert stats["items_removed"] == 1
    assert stats["items_modified"] == 1
    assert stats["items_unchanged"] == 1
    assert os.path.exists(res["output_diff_file"])


def test_boq_reconcile_and_anomalies():
    res_rec = boq_reconcile_with_field_reports("Curup", days_back=7)
    assert res_rec["status"] == "success"
    assert "reconciliation_summary" in res_rec

    res_anom = boq_check_project_anomalies(project_keyword="FTTH")
    assert res_anom["status"] == "success"
    assert "health_status" in res_anom


def test_generate_daily_executive_briefing():
    from alfa.tools.telecom_tools import generate_daily_executive_briefing
    res = generate_daily_executive_briefing(shift="morning", days_back=2, only_unreported=False, mark_as_reported=False)
    assert res["status"] == "success"
    assert "briefing_card" in res
    assert "DAILY EXECUTIVE BRIEFING" in res["briefing_card"]
    assert "PENGAMBILAN MATERIAL" in res["briefing_card"]
    assert "POINT PENTING" in res["briefing_card"]
    assert "by_cluster" in res["stats"]


def test_daily_briefing_anti_duplication_tracking():
    from alfa.core.db.wa_tracking import clear_wa_reported_state_sync
    from alfa.tools.telecom_tools import generate_daily_executive_briefing

    clear_wa_reported_state_sync()
    res1 = generate_daily_executive_briefing(shift="morning", days_back=2, only_unreported=True, mark_as_reported=True)
    assert res1["status"] == "success"

    # Second run should have 0 new messages because already reported
    res2 = generate_daily_executive_briefing(shift="morning", days_back=2, only_unreported=True, mark_as_reported=True)
    assert res2["status"] == "success"
    assert res2["stats"]["new_unreported_messages"] == 0
    assert "telah berhasil dilaporkan" in res2["briefing_card"]


def test_parse_delivery_order_document():
    from alfa.tools.telecom_tools import parse_delivery_order_document
    # Test on existing sample TOSBI file if exists
    sample_file = "/home/fahmial/Unduhan/TOSBI260425001.pdf"
    if os.path.exists(sample_file):
        res = parse_delivery_order_document(sample_file)
        assert res["status"] == "success"
        assert res["total_items_extracted"] >= 1
        assert res["delivery_order"]["number"] == "TOSBI260425001"


def test_wa_relay_send_from_telegram_validation():
    from alfa.tools.telecom_tools import wa_relay_send_from_telegram
    res = wa_relay_send_from_telegram("", "")
    assert res["status"] == "error"
    assert "tidak boleh kosong" in res["message"]


def test_dashboard_map_viewer_endpoints():
    from fastapi.testclient import TestClient

    from alfa.dashboard.app import create_app

    client = TestClient(create_app())
    res_map = client.get("/map")
    assert res_map.status_code == 200
    assert "ALFA FTTH Cluster Map Center" in res_map.text


def test_calculate_material_return_and_scrap():
    from alfa.tools.telecom_tools import calculate_material_return_and_scrap
    res = calculate_material_return_and_scrap(
        cluster_name="Curup",
        total_cable_issued_m=3000.0,
        total_poles_issued_pcs=40.0,
        cable_scrap_allowance_pct=3.0,
    )
    assert res["status"] == "success"
    assert "audit_card" in res
    assert "calculations" in res
    assert res["calculations"]["cable"]["allowed_scrap_m"] >= 0


def test_consolidate_delivery_orders_to_stock_ledger():
    from alfa.tools.telecom_tools import consolidate_delivery_orders_to_stock_ledger
    sample = "/home/fahmial/Unduhan/TOSBI260425001.pdf"
    if os.path.exists(sample):
        res = consolidate_delivery_orders_to_stock_ledger(file_paths=sample, output_filename="TEST_STOCK_LEDGER.xlsx")
        assert res["status"] == "success"
        assert res["total_material_types"] >= 1
        assert os.path.exists(res["output_file"])


def test_schedule_daily_briefing_task():
    from alfa.tools.telecom_tools import schedule_daily_briefing_task
    res = schedule_daily_briefing_task(enabled=True)
    assert res["status"] == "success"
    assert "job_id" in res


def test_transcribe_field_voice_note_file_not_found():
    from alfa.tools.telecom_tools import transcribe_field_voice_note
    res = transcribe_field_voice_note("/nonexistent/path/to/voice.ogg")
    assert res["status"] == "error"
    assert "tidak ditemukan" in res["message"]


def test_transcribe_field_voice_note_local_whisper_fallback(monkeypatch, tmp_path):
    from alfa.tools.telecom_tools import transcribe_field_voice_note
    # Create a dummy audio file
    dummy_audio = tmp_path / "voice_note.mp3"
    dummy_audio.write_bytes(b"dummy audio content")

    # Mock no Gemini API key to force local faster-whisper fallback
    monkeypatch.setattr(
        "alfa.swarm.llm_client.get_agent_api_client",
        lambda req: ("gemini", "", "gemini-2.5-flash", "", None),
    )

    # Mock _transcribe_local_whisper_fallback
    monkeypatch.setattr(
        "alfa.tools.telecom_tools._transcribe_local_whisper_fallback",
        lambda fpath, cluster_hint="": {
            "status": "success",
            "file": os.path.basename(fpath),
            "transcript": "Lapor bos penarikan kabel 150 meter dan 3 tiang selesai di cluster Melati.",
            "extracted_data": {
                "transcript": "Lapor bos penarikan kabel 150 meter dan 3 tiang selesai di cluster Melati.",
                "mandor": "",
                "cluster": cluster_hint or "Melati",
                "cables_meter": 150.0,
                "poles_count": 3,
                "fat_count": 0,
                "kendala": "",
            },
            "message": "Transkripsi voice note lapangan berhasil diekstrak via local faster-whisper (offline sovereign fallback).",
        },
    )

    res = transcribe_field_voice_note(str(dummy_audio), cluster_hint="Melati")
    assert res["status"] == "success"
    assert "Melati" in res["extracted_data"]["cluster"]
    assert res["extracted_data"]["cables_meter"] == 150.0
    assert res["extracted_data"]["poles_count"] == 3


def test_extract_telecom_quantities_strict_no_hallucination():
    from alfa.tools.telecom_tools import _extract_telecom_quantities_from_text

    # 1. Obrolan rencana/masa depan & nomor ID node tanpa satuan
    res1 = _extract_telecom_quantities_from_text("Besok gass fdt 2 pa 🔥🔥🔥")
    assert res1["fats"] == 0.0, "FDT 2 pada 'besok gass fdt 2 pa' adalah ID/rencana, bukan kuantitas 2 box!"
    assert res1["poles"] == 0.0
    assert res1["cables"] == 0.0

    # 2. Pengambilan tiang riil berformat shorthand resmi
    res2 = _extract_telecom_quantities_from_text("T72:10 btg\nT73:10 btg")
    assert res2["poles"] == 20.0
    assert res2["fats"] == 0.0

    # 3. Rencana tanam tiang (bukan realisasi)
    res3 = _extract_telecom_quantities_from_text("Rencana tanam tiang 5 btg besok pagi")
    assert res3["poles"] == 0.0, "Rencana masa depan tidak boleh dihitung sebagai realisasi tiang!"

    # 4. Nomor ID perangkat & port specification (bukan kuantitas barang)
    res4 = _extract_telecom_quantities_from_text("FDT 2 rusak mas tolong dicek")
    assert res4["fats"] == 0.0

    res5 = _extract_telecom_quantities_from_text("sudah sampai di FAT 16 port")
    assert res5["fats"] == 0.0

    # 5. Kuantitas riil dengan satuan eksplisit (box / unit)
    res6 = _extract_telecom_quantities_from_text("Selesai pasang FAT: 2 box hari ini")
    assert res6["fats"] == 2.0
    assert res6["fdts"] == 0.0

    # 6. Pemisahan ketat item FAT (akses) vs FDT (distribusi)
    res_fdt = _extract_telecom_quantities_from_text("Pengambilan material: FDT 96C: 2 unit dan FAT 16: 5 box")
    assert res_fdt["fdts"] == 2.0, "FDT harus tercatat 2 unit terpisah dari FAT!"
    assert res_fdt["fats"] == 5.0, "FAT harus tercatat 5 box terpisah dari FDT!"

    # 7. Kabel dengan konteks fisik
    res7 = _extract_telecom_quantities_from_text("Tarik kabel ADSS 500 mtr selesai")
    assert res7["cables"] == 500.0

    # 8. Bukan kabel (hanya ukuran jalan)
    res8 = _extract_telecom_quantities_from_text("lebar jalan 20 meter tidak bisa masuk mobil")
    assert res8["cables"] == 0.0


def test_briefing_does_not_count_plan_as_material_taking():
    from alfa.tools.telecom_tools import generate_daily_executive_briefing

    res = generate_daily_executive_briefing(shift="now", days_back=2, only_unreported=False, mark_as_reported=False)
    assert res["status"] == "success"
    card = res["briefing_card"]

    # Pastikan 'FAT/FDT' gabungan TIDAK muncul lagi, melainkan dipisah FAT dan FDT
    assert "Total FDT Terpasang" in card
    assert "Total FAT/ODP Terpasang" in card
    assert "Total FAT/FDT Terpasang" not in card, "FAT dan FDT tidak boleh digabung!"

    # Pastikan 'FAT/FDT: 2 box' TIDAK muncul pada grup pengambilan material lahat
    if "pengambilan material lahat" in card:
        lahat_section = card.split("pengambilan material lahat")[1].split("• 🏢")[0]
        assert "FAT: 2 box" not in lahat_section
        assert "FDT: 2 box" not in lahat_section
        assert "FAT/FDT" not in lahat_section
        # Bila muncul teks besok gass fdt 2 pa, harus di bawah label koordinasi/rencana, bukan laporan fisik
        if "Besok gass fdt 2 pa" in lahat_section:
            assert "💬 _Koordinasi/Rencana_" in lahat_section




